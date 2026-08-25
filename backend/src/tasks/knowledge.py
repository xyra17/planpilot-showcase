import csv
import io
import re
import uuid

from src.celery_app import celery_app
from src.core.time import utc_now
from src.tasks.runtime import run_async

MAX_EXTRACTED_CHARS = 50_000
TEXT_EXTENSIONS = {"txt", "md", "json", "pdf", "docx", "csv", "xlsx", "pptx"}
CHUNK_TARGET_CHARS = 1200
CHUNK_OVERLAP_CHARS = 200
EMBEDDING_BATCH_SIZE = 32
MAX_URL_RESPONSE_BYTES = 2 * 1024 * 1024
ALLOWED_URL_CONTENT_TYPES = {"application/json"}


class KnowledgeProcessingError(RuntimeError):
    pass


async def fetch_public_url_content(source_url: str) -> str:
    import httpx

    from src.services.ssrf_guard import UnsafeUrlError, normalize_public_http_url

    try:
        canonical_url = await normalize_public_http_url(source_url)
    except UnsafeUrlError as exc:
        raise KnowledgeProcessingError(str(exc)) from exc

    timeout = httpx.Timeout(connect=5.0, read=20.0, write=5.0, pool=5.0)
    try:
        async with httpx.AsyncClient(
            timeout=timeout,
            follow_redirects=False,
            trust_env=False,
        ) as client:
            async with client.stream(
                "GET",
                f"https://r.jina.ai/{canonical_url}",
                headers={"Accept": "text/plain", "User-Agent": "PlanPilot-URL-Importer/1.0"},
            ) as response:
                if 300 <= response.status_code < 400:
                    raise KnowledgeProcessingError("网址抓取不允许重定向")
                response.raise_for_status()
                content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
                if content_type and not (
                    content_type.startswith("text/") or content_type in ALLOWED_URL_CONTENT_TYPES
                ):
                    raise KnowledgeProcessingError("网址返回了不支持的内容类型")
                declared_size = response.headers.get("content-length")
                if declared_size and int(declared_size) > MAX_URL_RESPONSE_BYTES:
                    raise KnowledgeProcessingError("网址内容超过 2 MB 限制")
                chunks: list[bytes] = []
                received = 0
                async for chunk in response.aiter_bytes():
                    received += len(chunk)
                    if received > MAX_URL_RESPONSE_BYTES:
                        raise KnowledgeProcessingError("网址内容超过 2 MB 限制")
                    chunks.append(chunk)
    except KnowledgeProcessingError:
        raise
    except (httpx.HTTPError, ValueError) as exc:
        raise KnowledgeProcessingError("网址内容抓取失败") from exc
    return b"".join(chunks).decode("utf-8", errors="replace")[:MAX_EXTRACTED_CHARS].strip()


def chunk_text(
    content: str,
    target_chars: int = CHUNK_TARGET_CHARS,
    overlap_chars: int = CHUNK_OVERLAP_CHARS,
) -> list[tuple[str, int, int]]:
    """Split text near paragraph/sentence boundaries while retaining source offsets."""
    normalized = content.strip()
    if not normalized:
        return []
    if target_chars <= 0 or overlap_chars < 0 or overlap_chars >= target_chars:
        raise ValueError("分块参数无效")

    chunks: list[tuple[str, int, int]] = []
    start = 0
    length = len(normalized)
    boundary_pattern = re.compile(r"\n\s*\n|(?<=[。！？.!?])\s+|\n")

    while start < length:
        ideal_end = min(start + target_chars, length)
        end = ideal_end
        if ideal_end < length:
            search_start = min(start + target_chars // 2, ideal_end)
            matches = list(boundary_pattern.finditer(normalized, search_start, ideal_end))
            if matches:
                end = matches[-1].end()

        chunk_start = start
        chunk_end = end
        while chunk_start < chunk_end and normalized[chunk_start].isspace():
            chunk_start += 1
        while chunk_end > chunk_start and normalized[chunk_end - 1].isspace():
            chunk_end -= 1
        if chunk_start < chunk_end:
            chunks.append((normalized[chunk_start:chunk_end], chunk_start, chunk_end))

        if end >= length:
            break
        next_start = max(end - overlap_chars, start + 1)
        while next_start < end and normalized[next_start].isspace():
            next_start += 1
        start = next_start

    return chunks


async def _embed_chunks(
    clients: list,
    title: str,
    chunks: list[tuple[str, int, int]],
    model: str = "qwen3-embedding-0.6b",
    dimensions: int = 1024,
) -> list[list[float]]:
    last_error: Exception | None = None
    inputs = [f"{title}\n{chunk}" for chunk, _, _ in chunks]
    for client in clients:
        try:
            embeddings: list[list[float]] = []
            for offset in range(0, len(inputs), EMBEDDING_BATCH_SIZE):
                response = await client.embeddings.create(
                    model=model,
                    input=inputs[offset : offset + EMBEDDING_BATCH_SIZE],
                )
                batch = [row.embedding for row in response.data]
                if len(batch) != len(inputs[offset : offset + EMBEDDING_BATCH_SIZE]):
                    raise KnowledgeProcessingError("Embedding 返回数量与分块数量不一致")
                if any(len(vector) != dimensions for vector in batch):
                    raise KnowledgeProcessingError(f"Embedding 维度错误：期望 {dimensions}")
                embeddings.extend(batch)
            return embeddings
        except Exception as exc:
            last_error = exc
    raise KnowledgeProcessingError("所有 Embedding 服务均调用失败") from last_error


def extract_text(raw: bytes, ext: str) -> str:
    """Extract bounded text and raise a useful error instead of silently succeeding."""
    try:
        if ext in {"txt", "md", "json"}:
            content = raw.decode("utf-8", errors="replace")
        elif ext == "pdf":
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise KnowledgeProcessingError("PDF 已加密，无法解析")
            content = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == "docx":
            import docx

            document = docx.Document(io.BytesIO(raw))
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        elif ext == "csv":
            reader = csv.reader(io.StringIO(raw.decode("utf-8", errors="replace")))
            content = "\n".join(",".join(row) for row in reader)
        elif ext == "xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines: list[str] = []
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    lines.append("\t".join("" if value is None else str(value) for value in row))
            content = "\n".join(lines)
        elif ext == "pptx":
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(raw))
            slides: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                fragments = [f"幻灯片 {index}"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        fragments.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            fragments.append("\t".join(cell.text for cell in row.cells))
                slides.append("\n".join(fragments))
            content = "\n\n".join(slides)
        else:
            raise KnowledgeProcessingError("该文件仅作为附件保存，不需要建立索引")
    except KnowledgeProcessingError:
        raise
    except Exception as exc:
        raise KnowledgeProcessingError(f"文件解析失败：{type(exc).__name__}") from exc

    content = content.strip()
    if not content:
        raise KnowledgeProcessingError("未提取到文本；扫描版 PDF 请先进行 OCR")
    return content[:MAX_EXTRACTED_CHARS]


@celery_app.task(
    name="src.tasks.knowledge.process_knowledge_item",
    bind=True,
    max_retries=3,
    default_retry_delay=60,
)
def process_knowledge_item(self, item_id: str):
    try:
        run_async(_process(item_id))
    except Exception as exc:
        if self.request.retries >= self.max_retries:
            run_async(_mark_failed(item_id, exc, self.request.retries))
            raise
        run_async(_mark_retry(item_id, exc, self.request.retries + 1))
        raise self.retry(exc=exc)


# Backwards-compatible task name for workers or queued messages from the previous version.
@celery_app.task(name="src.tasks.knowledge.vectorize_item")
def vectorize_item(item_id: str):
    process_knowledge_item.delay(item_id)


async def _process(item_id: str) -> None:
    from openai import AsyncOpenAI
    from sqlalchemy import delete, select
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from src.config import settings
    from src.models import KnowledgeChunk, KnowledgeItem
    from src.services.runtime_model_config import get_runtime_model_config
    from src.services.object_storage import (
        ObjectStorageError,
        extension_for_reference,
        get_object_storage,
    )

    engine = create_async_engine(settings.database_url)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with session_factory() as db:
            item = (
                await db.execute(select(KnowledgeItem).where(KnowledgeItem.id == item_id))
            ).scalar_one_or_none()
            if not item:
                return

            item.processing_error = None
            ext = extension_for_reference(item.file_path) if item.file_path else ""

            if item.source_type == "url" and item.source_url:
                item.processing_status = "parsing"
                await db.commit()
                item.content = await fetch_public_url_content(item.source_url)
                if not item.content:
                    raise KnowledgeProcessingError("页面内容为空，无法导入")
                item.content_length = len(item.content)
                item.content_format = "markdown" if ext == "md" else "plain"
                for line in item.content.splitlines():
                    if line.strip():
                        item.title = line.strip()[:120]
                        break
            elif item.file_path and ext in TEXT_EXTENSIONS:
                item.processing_status = "parsing"
                await db.commit()
                try:
                    item.content = extract_text(
                        await get_object_storage().read(item.file_path), ext
                    )
                except ObjectStorageError as exc:
                    raise KnowledgeProcessingError("原始文件不存在") from exc
                item.content_length = len(item.content)
            elif not item.content:
                # Binary attachments are usable without semantic indexing.
                item.processing_status = "ready"
                item.processed_at = utc_now()
                await db.commit()
                return

            chunks = chunk_text(item.content)
            if not chunks:
                raise KnowledgeProcessingError("文档内容为空，无法分块")

            item.processing_status = "embedding"
            await db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.item_id == item.id))
            await db.commit()

            runtime_models = get_runtime_model_config()
            clients = []
            if runtime_models.embedding_enabled and runtime_models.embedding_base_url:
                clients.append(
                    AsyncOpenAI(
                        api_key=runtime_models.embedding_api_key or "local",
                        base_url=runtime_models.embedding_base_url,
                        timeout=settings.embedding_timeout_seconds,
                        max_retries=settings.embedding_max_retries,
                    )
                )
            if not clients:
                raise KnowledgeProcessingError("未配置可用的 Embedding 服务")

            embeddings = await _embed_chunks(
                clients,
                item.title,
                chunks,
                model=runtime_models.embedding_model_name,
                dimensions=runtime_models.embedding_dimensions,
            )
            for index, ((chunk, start_char, end_char), embedding) in enumerate(
                zip(chunks, embeddings, strict=True)
            ):
                db.add(
                    KnowledgeChunk(
                        id=str(uuid.uuid4()),
                        item_id=item.id,
                        chunk_index=index,
                        content=chunk,
                        start_char=start_char,
                        end_char=end_char,
                        embedding=embedding,
                    )
                )

            # Retain a document-level vector for backwards compatibility.
            item.embedding = embeddings[0]
            item.processing_status = "ready"
            item.processing_error = None
            item.processed_at = utc_now()
            await db.commit()
    finally:
        await engine.dispose()


async def _mark_retry(item_id: str, exc: Exception, retry_count: int) -> None:
    from sqlalchemy import select
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from src.config import settings
    from src.models import KnowledgeItem

    engine = create_async_engine(settings.database_url)
    try:
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            item = (
                await db.execute(select(KnowledgeItem).where(KnowledgeItem.id == item_id))
            ).scalar_one_or_none()
            if item:
                item.processing_status = "queued"
                item.processing_error = str(exc)[:500]
                item.retry_count = retry_count
                await db.commit()
    finally:
        await engine.dispose()


async def _mark_failed(item_id: str, exc: Exception, retry_count: int) -> None:
    from sqlalchemy import select
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from src.config import settings
    from src.models import KnowledgeItem

    engine = create_async_engine(settings.database_url)
    try:
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            item = (
                await db.execute(select(KnowledgeItem).where(KnowledgeItem.id == item_id))
            ).scalar_one_or_none()
            if item:
                item.processing_status = "failed"
                item.processing_error = str(exc)[:500]
                item.retry_count = retry_count
                await db.commit()
    finally:
        await engine.dispose()
